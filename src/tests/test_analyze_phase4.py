"""Unit tests for the Phase-4 analysis toolkit (scripts/analyze_phase4.py).

Uses a small synthetic results tree (2 datasets x 1 subset x 10 methods x
2 seeds + deterministic rows) with generated npz artifacts under tmp_path,
so no real sweep data is touched. Covers: pooling math, tie/TOST machinery,
hypothesis + prediction subcommands end to end, manifest planned==done,
npz recompute verification, and the reviewed XLSX fill.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

SRC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SRC / "scripts"))

import analyze_phase4 as A  # noqa: E402

SEEDS = [42, 123]
SEEDED = ["c_qrnn", "c_mdn", "uacqr", "distcp", "weighted_cqr", "wqlcp",
          "edl", "cxgbq"]


def _make_npz(path: Path, break_unit2: bool = False) -> None:
    y = np.arange(1, 21, dtype=np.float64)
    lo = y - 5.0
    up = y + 5.0
    if break_unit2:
        up[10:] = y[10:] - 2.0  # unit 2 fully uncovered
    units = np.array(["U1"] * 10 + ["U2"] * 10, dtype=np.str_)
    stage = np.array([0] * 7 + [1] * 6 + [2] * 7, dtype=np.int64)
    np.savez(path, y_true=y, lower=lo.astype(np.float32),
             median=y.astype(np.float32), upper=up.astype(np.float32),
             units=units, stage=stage)


def _cell(ds, subset, method, seed, n_test=20, npz: Path | None = None):
    u1, u2 = f"{subset}_1", f"{subset}_2"
    # metrics consistent with the npz (as the real runner records them)
    if npz is not None:
        z = np.load(npz)
        yt, lo_, up_, med = z["y_true"], z["lower"], z["upper"], z["median"]
        picp = A.M.calculate_picp(yt, lo_, up_)
        nmpiw = A.M.calculate_nmpiw(yt, lo_, up_)
        cwc = A.M.calculate_cwc(yt, lo_, up_)
        wink = A.M.winkler_score(yt, lo_, up_)
        mae = float(np.mean(np.abs(yt - med)))
        rmse = float(np.sqrt(np.mean((yt - med) ** 2)))
    else:
        picp, nmpiw, cwc, wink, mae, rmse = 1.0, 0.95, 0.95, 10.0, 0.0, 0.0
    pu_picp = [1.0, 0.0] if picp < 0.99 else [1.0, 1.0]
    return {
        "dataset": ds, "subset": subset, "method": method, "seed": seed,
        "status": "OK", "reason": None, "q_hat": 5.0, "PICP": picp,
        "RS": abs(picp - 0.9), "NMPIW": nmpiw, "CWC": cwc, "Winkler": wink,
        "MAE": mae, "RMSE": rmse, "R2": 1.0, "RUL_score": 0.0,
        "crossing_freq": 0.0, "crossing_note": None,
        "per_unit": {"units": [u1, u2], "picp": pu_picp, "n": [10, 10],
                     "mean_picp": float(np.mean(pu_picp))},
        "per_stage": {s: {"PICP": picp, "n": n} for s, n in
                      (("early", 7), ("mid", 6), ("late", 7))},
        "diagnostics": {"ess_fraction": 0.8, "clip_rate": 0.01,
                        "collapsed": False, "fell_back": False},
        "efficiency": {"footprint_mb": 0.1, "latency_ms_per_sample": 0.001},
        "time_s": 1.0, "n_train": 100, "n_val": 10, "n_cal": 20,
        "n_test": n_test, "n_units_train": 4, "n_units_test": 2,
        "pred_path": "",
    }


@pytest.fixture()
def world(tmp_path, monkeypatch):
    """Synthetic results tree + npz artifacts; PRED_DIR patched to tmp."""
    pred = tmp_path / "pred"
    pred.mkdir()
    monkeypatch.setattr(A, "PRED_DIR", pred)

    def build():
        results = {"meta": {}, "ncmapss": {}, "cmapss": {}, "arch_ablation": {}}
        for ds, subset in (("ncmapss", "DS01"), ("cmapss", "FD001")):
            results[ds][subset] = {}
            for m in SEEDED:
                cells = []
                for sd in SEEDS:
                    name = f"{ds}_{subset}_{m}_s{sd}.npz"
                    _make_npz(pred / name)
                    c = _cell(ds, subset, m, sd, npz=pred / name)
                    c["pred_path"] = f"pred/{name}"
                    cells.append(c)
                results[ds][subset][m] = cells
            for m in ("cgpr", "raw_qrnn"):
                name = f"{ds}_{subset}_{m}.npz"
                _make_npz(pred / name)
                c = _cell(ds, subset, m, None, npz=pred / name)
                c["pred_path"] = f"pred/{name}"
                results[ds][subset][m] = [c]
        # one broken cell: c_mdn seed 42 on DS01 has unit-2 uncovered
        bad = pred / "ncmapss_DS01_c_mdn_s42.npz"
        _make_npz(bad, break_unit2=True)
        c = _cell("ncmapss", "DS01", "c_mdn", 42, npz=bad)
        c["pred_path"] = "pred/ncmapss_DS01_c_mdn_s42.npz"
        results["ncmapss"]["DS01"]["c_mdn"][0] = c
        # one FAIL cell for the inventory
        fail = _cell("ncmapss", "DS01", "wqlcp", 999)
        fail["status"] = "FAIL"
        fail["reason"] = "synthetic failure"
        results["ncmapss"]["DS01"]["wqlcp"].append(fail)
        # ablation
        results["arch_ablation"]["DS01"] = {
            "width_32": [_cell("ncmapss", "DS01", "arch_ablation", 42)],
        }
        results["arch_ablation"]["DS01"]["width_32"][0]["config"] = {
            "id": "width_32", "factor": "width", "value": 32}
        return results

    matrix = {"schema_version": 1, "split_seed": 42, "seeds": SEEDS, "jobs": [
        {"dataset": "ncmapss", "subset": "DS01", "method": "c_qrnn",
         "seed": 42},
        {"dataset": "ncmapss", "subset": "DS01", "method": "c_qrnn",
         "seed": 123},
        {"dataset": "cmapss", "subset": "FD001", "method": "no_such",
         "seed": 42},  # missing on purpose
    ]}
    matrix_path = tmp_path / "run_matrix.json"
    matrix_path.write_text(json.dumps(matrix))
    return {"results": build(), "matrix": matrix_path, "tmp": tmp_path}


# ── pooling & table smoke ─────────────────────────────────────────────────


class TestPoolingAndTables:
    def test_pooled_coverage_is_sample_weighted(self):
        c1 = {"PICP": 0.8, "n_test": 10}
        c2 = {"PICP": 0.9, "n_test": 30}
        assert A.pooled_coverage([c1, c2]) == pytest.approx(
            (0.8 * 10 + 0.9 * 30) / 40)

    def test_per_subset_dev_table_columns_align(self, world):
        # regression: the |PICP-0.90| rows were built as [sub] + row_d[1:],
        # dropping the first method's column in every row (2026-09-24)
        out = A.t_per_subset(world["results"])
        tables, cur = [], []
        for line in out.splitlines():
            if line.startswith("|") and line.endswith("|"):
                cur.append(line)
            elif cur:
                tables.append(cur)
                cur = []
        if cur:
            tables.append(cur)
        assert tables
        for t in tables:
            counts = {len(r.strip().strip("|").split("|")) for r in t}
            assert len(counts) == 1, "misaligned table:\n" + "\n".join(t[:5])

    def test_aggregate_table_lists_all_methods(self, world):
        md = A.t_aggregate(world["results"])
        for lab in A.METHOD_LABEL.values():
            assert lab in md

    def test_tables_smoke(self, world):
        for fn in (A.t_per_seed, A.t_per_subset, A.t_per_unit, A.t_per_stage,
                   A.t_crossing, A.t_shift, A.t_ablation):
            out = fn(world["results"])
            assert out and "|" in out  # markdown table present


# ── tie / TOST machinery ──────────────────────────────────────────────────


class TestTieMachinery:
    def _frames(self, shift: float = 0.0):
        idx = pd.Index([f"u{i}" for i in range(20)], name="unit")
        base = pd.DataFrame({"value": np.linspace(0.85, 0.95, 20)}, index=idx)
        shifted = base + shift
        return {m: base.copy() for m in A.SPLIT_CP_4} | {
            "distcp": shifted.copy()}

    def test_perfect_tie_is_equivalent(self):
        res = A.four_arm_tie(self._frames(0.0))
        assert res["tie"] is True
        assert res["n_units"] == 20

    def test_large_shift_breaks_tie(self):
        res = A.four_arm_tie(self._frames(0.2))
        assert res["tie"] is False

    def test_paired_bootstrap_ci_brackets_mean(self):
        idx = pd.Index([f"u{i}" for i in range(50)], name="unit")
        d = pd.Series(np.full(50, 0.05), index=idx, name="diff")
        mean, lo, hi = A.paired_bootstrap_ci(d, n_boot=200)
        # constant differences -> degenerate CI at the mean (allow fp slack)
        assert mean == pytest.approx(0.05)
        assert lo <= 0.05 + 1e-12 and hi >= 0.05 - 1e-12

    def test_tost_power_mde_scales_with_n(self):
        small = A.tost_power_mde(10, 0.1, 0.03)
        big = A.tost_power_mde(100, 0.1, 0.03)
        assert big < small


# ── hypothesis subcommands (end to end on synthetic data) ─────────────────


class TestHypotheses:
    def test_h1_runs(self, world):
        md, st = A.h1_shift_validity(world["results"])
        assert "H1" in md and st["verdict"]

    def test_h2_runs(self, world):
        md, st = A.h2_local_adaptivity(world["results"])
        assert "H2" in md and st["verdict"]

    def test_h3_runs(self, world):
        md, st = A.h3_fusion_vs_series(world["results"])
        assert "H3" in md and st["verdict"]

    def test_h4_runs(self, world):
        md, st = A.h4_baselines(world["results"])
        assert "H4" in md and st["verdict"]


# ── prediction subcommand ─────────────────────────────────────────────────


class TestPrediction:
    def test_eleven_rows_and_hit_rate(self, world):
        md, st = A.build_prediction_report(world["results"])
        assert len(st["rows"]) == 11
        assert st["n"] == 11
        assert 0 <= st["hits"] <= 11
        lo, hi = st["ci"]
        assert 0.0 <= lo <= st["hits"] / 11 <= hi <= 1.0
        assert "P11" in md and "Hit rate" in md

    def test_fill_xlsx_only_touches_efgh(self, world):
        import openpyxl
        xlsx = world["tmp"] / "tracker.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Predictions"
        ws.append(["ID", "Metric", "Predicted best", "Rationale",
                   "Status", "Observed best", "Why (if MISS)",
                   "Date recorded"])
        for i in range(1, 12):
            ws.append([f"P{i}", f"metric {i}", f"pred {i}", f"rat {i}",
                       "PENDING", None, None, None])
        wb.save(xlsx)

        spec = world["tmp"] / "spec.json"
        spec.write_text(json.dumps({
            "P1": {"status": "HIT", "observed": "Weighted-CP CQR",
                   "why": None, "date": "2026-09-24"},
            "P2": {"status": "MISS", "observed": "tie broken: C-QRNN",
                   "why": "implementation effect", "date": "2026-09-24"},
        }))
        out = A.fill_xlsx(xlsx, spec)
        assert "PENDING remaining: 9" in out
        wb2 = openpyxl.load_workbook(xlsx)
        ws2 = wb2["Predictions"]
        # filled rows
        assert ws2.cell(row=2, column=5).value == "HIT"
        assert ws2.cell(row=2, column=6).value == "Weighted-CP CQR"
        assert ws2.cell(row=2, column=8).value == "2026-09-24"
        assert ws2.cell(row=3, column=7).value == "implementation effect"
        # frozen columns untouched
        for r in range(2, 13):
            assert ws2.cell(row=r, column=3).value.startswith("pred")
            assert ws2.cell(row=r, column=4).value.startswith("rat")
        # untouched rows still PENDING
        assert ws2.cell(row=4, column=5).value == "PENDING"


# ── E5 verification ───────────────────────────────────────────────────────


class TestVerify:
    def test_manifest_missing_detected(self, world):
        md, st = A.verify_manifest(world["results"], world["matrix"])
        assert st["n_jobs"] == 3
        assert st["n_missing"] == 1  # the no_such job
        assert "MISSING" in md

    def test_manifest_counts_ablation_cell_done(self, world):
        # regression: iter_cells must walk the top-level arch_ablation
        # section, or the 18 ablation manifest jobs read as missing even
        # though their cells are present (caught in the full sweep, 2026-09-24)
        m = json.loads(world["matrix"].read_text())
        m["jobs"].append({"dataset": "ncmapss", "subset": "DS01",
                          "method": "arch_ablation", "seed": 42,
                          "config": {"id": "width_32", "factor": "width",
                                     "value": 32}})
        p = world["tmp"] / "matrix_with_ablation.json"
        p.write_text(json.dumps(m))
        md, st = A.verify_manifest(world["results"], p)
        assert st["n_jobs"] == 4
        assert st["n_missing"] == 1  # still only no_such

    def test_recompute_matches_npz(self, world):
        md, st = A.verify_recompute(world["results"])
        assert st["n_checked"] > 0
        assert st["n_fail"] == 0  # JSON values consistent with npz

    def test_fail_inventory(self, world):
        md, st = A.verify_fail_inventory(world["results"])
        assert st["n_fail_cells"] == 1
        assert "synthetic failure" in md

    def test_patch_winkler_repairs_inflated_cell(self, world):
        res = world["results"]
        # select by key (seed), never list position
        cell = next(c for c in res["ncmapss"]["DS01"]["c_qrnn"]
                    if c["seed"] == 42)
        cell["Winkler"] = cell["Winkler"] + 10.0  # mimic buggy runner
        p = world["tmp"] / "res.json"
        p.write_text(json.dumps(res))
        md, st = A.patch_winkler(json.loads(p.read_text()), p)
        assert st["n_fixed"] == 1
        reloaded = json.loads(p.read_text())
        patched = next(c for c in reloaded["ncmapss"]["DS01"]["c_qrnn"]
                       if c["seed"] == 42)
        z = np.load(world["tmp"] / patched["pred_path"])
        assert patched["Winkler"] == pytest.approx(
            A.M.winkler_score(z["y_true"], z["lower"], z["upper"]))
        assert "winkler_patch" in reloaded["meta"]

    def test_patch_winkler_covers_ablation_cell(self, world):
        # regression: patch_winkler iterates via iter_cells; it must reach the
        # arch_ablation section too (ablation cells are stored there, not under
        # results[dataset][subset])
        res = world["results"]
        cell = res["arch_ablation"]["DS01"]["width_32"][0]
        npz = world["tmp"] / "pred" / "ncmapss_DS01_arch_ablation_s42.npz"
        _make_npz(npz)
        cell["pred_path"] = "pred/ncmapss_DS01_arch_ablation_s42.npz"
        cell["Winkler"] = cell["Winkler"] + 10.0  # mimic buggy runner
        p = world["tmp"] / "res.json"
        p.write_text(json.dumps(res))
        md, st = A.patch_winkler(json.loads(p.read_text()), p)
        assert st["n_fixed"] == 1
        reloaded = json.loads(p.read_text())
        patched = reloaded["arch_ablation"]["DS01"]["width_32"][0]
        z = np.load(world["tmp"] / patched["pred_path"])
        assert patched["Winkler"] == pytest.approx(
            A.M.winkler_score(z["y_true"], z["lower"], z["upper"]))


class TestPredictionFixes:
    def test_p7_condcov_differs_by_method(self, world):
        # regression: _p7_condcov must filter cells by method; without the
        # filter every method scored identically (mean over ALL cells)
        res = {"meta": {}, "ncmapss": {}, "cmapss": {}}
        res["ncmapss"]["DS01"] = {
            "uacqr": [_cell("ncmapss", "DS01", "uacqr", 42)],
            "c_qrnn": [_cell("ncmapss", "DS01", "c_qrnn", 42)],
        }
        u, q = res["ncmapss"]["DS01"]["uacqr"][0], \
            res["ncmapss"]["DS01"]["c_qrnn"][0]
        u["PICP"] = 0.90
        u["per_stage"] = {s: {"PICP": 0.90, "n": 7}
                          for s in ("early", "mid", "late")}
        q["PICP"] = 0.80
        q["per_stage"] = {s: {"PICP": 0.80, "n": 7}
                          for s in ("early", "mid", "late")}
        winners, out = A._p7_condcov(res)
        assert out["uacqr"] == pytest.approx(0.0)
        assert out["c_qrnn"] == pytest.approx(0.1)
        assert winners == ["uacqr"]
